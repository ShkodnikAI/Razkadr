"""Build the reference Razkadr PPTX template (naryad 022).

Writes ``templates/razkadr-default.pptx``: a 16:9 presentation with exactly
five slide layouts named after the layout engine's macro layouts (naryad 021)
- ``title``, ``hero_caption``, ``ingredients_list``, ``steps_grid``,
``serving`` - each with named placeholders matching the slides.json slot
names (``heading``, ``subheading``, ``frame``, ``caption``, ``timecode``,
``items`` and ``cell1..6_{number,frame,text,timecode}``).

Project palette (docs/brand): burgundy #6B0F1A, champagne #D4AF7F,
cream #F8F4E9, graphite #2B2B2B. Font: DejaVu Sans - shipped with
LibreOffice/desktop Linux CI and covers Cyrillic and Latin.

The build is deterministic: the same script version always produces the
same package parts (zip entry timestamps differ, contents do not - the
reproducibility test compares decompressed part bytes).

Usage: python scripts/templates/build_default.py [--out PATH]
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu

TEMPLATE_RELPATH = "templates/razkadr-default.pptx"

# palette
BURGUNDY = "6B0F1A"
CHAMPAGNE = "D4AF7F"
CREAM = "F8F4E9"
GRAPHITE = "2B2B2B"
FONT = "DejaVu Sans"

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5

IN = 914400  # EMU per inch


def _emu(inches: float) -> int:
    return round(inches * IN)


# --- placeholder XML builders -------------------------------------------------

NSDECL = (
    'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)

# bodyPr insets: 0.06in sides, 0.03in top/bottom (kept in sync with the
# renderer's text metrics in src/razkadr/render/pptx.py)
L_INS = _emu(0.06)
T_INS = _emu(0.03)


def ph_text(
    phid: int,
    name: str,
    *,
    ph_type: str | None,
    idx: int | None,
    x: float,
    y: float,
    w: float,
    h: float,
    sz: int,
    color: str = GRAPHITE,
    bold: bool = False,
    align: str = "l",
    anchor: str = "t",
    bullet: bool = False,
) -> str:
    """A text placeholder ``p:sp`` for a slide layout.

    ``sz`` is the start size in pt; the renderer's auto-fit may shrink it
    down to the floor (naryad 022). Bullet: literal ``"• "`` prefixes are
    rendered by the renderer (metrics stay exact), so every placeholder
    sets ``<a:buNone/>``.
    """
    type_attr = f' type="{ph_type}"' if ph_type else ""
    idx_attr = f' idx="{idx}"' if idx is not None else ""
    return (
        f'<p:sp {NSDECL}>'
        f"<p:nvSpPr>"
        f'<p:cNvPr id="{phid}" name="{name}"/>'
        f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
        f"<p:nvPr><p:ph{type_attr}{idx_attr}/></p:nvPr>"
        f"</p:nvSpPr>"
        f"<p:spPr>"
        f'<a:xfrm><a:off x="{_emu(x)}" y="{_emu(y)}"/><a:ext cx="{_emu(w)}" cy="{_emu(h)}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f"</p:spPr>"
        f"<p:txBody>"
        f'<a:bodyPr wrap="square" lIns="{L_INS}" tIns="{T_INS}" rIns="{L_INS}" bIns="{T_INS}" anchor="{anchor}">'
        f"<a:normAutofit/></a:bodyPr>"
        f"<a:lstStyle>"
        f'<a:lvl1pPr algn="{align}" marL="0" indent="0">'
        f"<a:buNone/>"
        f'<a:defRPr sz="{sz * 100}" b="{1 if bold else 0}">'
        f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
        f'<a:latin typeface="{FONT}"/><a:cs typeface="{FONT}"/>'
        f"</a:defRPr>"
        f"</a:lvl1pPr>"
        f"</a:lstStyle>"
        f"<a:p/>"
        f"</p:txBody>"
        f"</p:sp>"
    )


def ph_pic(phid: int, name: str, idx: int, x: float, y: float, w: float, h: float) -> str:
    """A picture placeholder (``p:sp`` with ``type="pic"``, same shape the
    built-in layouts use). The renderer reads its geometry, removes the
    empty placeholder and inserts the real frame fitted inside the box
    with proportions preserved."""
    return (
        f'<p:sp {NSDECL}>'
        f"<p:nvSpPr>"
        f'<p:cNvPr id="{phid}" name="{name}"/>'
        f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
        f'<p:nvPr><p:ph type="pic" idx="{idx}"/></p:nvPr>'
        f"</p:nvSpPr>"
        f"<p:spPr>"
        f'<a:xfrm><a:off x="{_emu(x)}" y="{_emu(y)}"/><a:ext cx="{_emu(w)}" cy="{_emu(h)}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f"</p:spPr>"
        f"<p:txBody>"
        f"<a:bodyPr/>"
        f"<a:lstStyle/>"
        f"<a:p/>"
        f"</p:txBody>"
        f"</p:sp>"
    )


def static_text(phid: int, name: str, x: float, y: float, w: float, h: float, sz: int, color: str, bold: bool = False) -> str:
    """A STATIC text shape baked into the layout: no ``<p:ph/>``, so it is
    never cloned onto slides and never participates in placeholder
    inheritance (the brand footer)."""
    return (
        f'<p:sp {NSDECL}>'
        f"<p:nvSpPr>"
        f'<p:cNvPr id="{phid}" name="{name}"/>'
        f"<p:cNvSpPr/>"
        f"<p:nvPr/>"
        f"</p:nvSpPr>"
        f"<p:spPr>"
        f'<a:xfrm><a:off x="{_emu(x)}" y="{_emu(y)}"/><a:ext cx="{_emu(w)}" cy="{_emu(h)}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f'<a:noFill/>'
        f"</p:spPr>"
        f"<p:txBody>"
        f'<a:bodyPr wrap="square" anchor="ctr"/>'
        f"<a:lstStyle/>"
        f'<a:p><a:r><a:rPr lang="en-US" sz="{sz * 100}" b="{1 if bold else 0}">'
        f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
        f'<a:latin typeface="{FONT}"/><a:cs typeface="{FONT}"/>'
        f"</a:rPr><a:t>Razkadr</a:t></a:r></a:p>"
        f"</p:txBody>"
        f"</p:sp>"
    )


def deco_rect(phid: int, x: float, y: float, w: float, h: float, color: str) -> str:
    """A static decorative rectangle baked into the layout (not a
    placeholder; never cloned onto slides)."""
    return (
        f'<p:sp {NSDECL}>'
        f"<p:nvSpPr>"
        f'<p:cNvPr id="{phid}" name="deco-{phid}"/>'
        f"<p:cNvSpPr/>"
        f"<p:nvPr/>"
        f"</p:nvSpPr>"
        f"<p:spPr>"
        f'<a:xfrm><a:off x="{_emu(x)}" y="{_emu(y)}"/><a:ext cx="{_emu(w)}" cy="{_emu(h)}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
        f"<a:ln><a:noFill/></a:ln>"
        f"</p:spPr>"
        f"<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody>"
        f"</p:sp>"
    )


# --- layout plans ---------------------------------------------------------------

# steps_grid cell geometry (nominal 3x2 arrangement; the renderer recomputes
# the grid per slides.json "columns" - 2 -> 2x2, 3 -> 3x2). Grid area below
# the heading.
GRID_X, GRID_Y = 0.5, 1.35
GRID_W, GRID_H = SLIDE_W_IN - 1.0, SLIDE_H_IN - 1.35 - 0.35
CELL_W, CELL_H = GRID_W / 3, GRID_H / 2

# frame sub-rect inside a cell: left part of the cell below the badge row
BADGE_H = 0.38
FRAME_W_FRAC = 0.46
FRAME_PAD = 0.12


def cell_specs() -> list[dict]:
    """6 cell groups (cell1..cell6) in a nominal 3x2 arrangement."""
    specs = []
    for i in range(1, 7):
        col, row = (i - 1) % 3, (i - 1) // 3
        cx, cy = GRID_X + col * CELL_W, GRID_Y + row * CELL_H
        specs.append(
            {
                "number": (cx, cy, 0.55, BADGE_H),
                "timecode": (cx + CELL_W - 1.3, cy, 1.3, BADGE_H),
                "frame": (cx, cy + BADGE_H, CELL_W * FRAME_W_FRAC, CELL_H - BADGE_H - FRAME_PAD),
                "text": (
                    cx + CELL_W * FRAME_W_FRAC + FRAME_PAD,
                    cy + BADGE_H,
                    CELL_W * (1 - FRAME_W_FRAC) - FRAME_PAD,
                    CELL_H - BADGE_H - FRAME_PAD,
                ),
                "origin": (cx, cy),
            }
        )
    return specs


def layout_bodies() -> dict[str, list[str]]:
    """XML of the shapes (placeholders + decor) per layout, in z-order."""
    bodies: dict[str, list[str]] = {}
    phid = iter(range(2, 200))  # shape ids are layout-local and unique

    # title: burgundy bar left, heading + subheading left, frame right half
    bodies["title"] = [
        deco_rect(next(phid), 0, 0, 0.18, SLIDE_H_IN, BURGUNDY),
        ph_text(
            next(phid), "heading", ph_type="ctrTitle", idx=None,
            x=0.7, y=1.5, w=6.35, h=2.3, sz=40, color=BURGUNDY, bold=True, align="l",
        ),
        ph_text(
            next(phid), "subheading", ph_type="subTitle", idx=1,
            x=0.7, y=3.9, w=6.1, h=1.9, sz=20, color=GRAPHITE, align="l",
        ),
        ph_pic(next(phid), "frame", 10, 7.4, 0, SLIDE_W_IN - 7.4, SLIDE_H_IN),
    ]

    # hero_caption: full-width frame on top, caption below, timecode bottom-right
    bodies["hero_caption"] = [
        deco_rect(next(phid), 0, 0, SLIDE_W_IN, 0.14, BURGUNDY),
        ph_pic(next(phid), "frame", 10, 0.5, 0.42, SLIDE_W_IN - 1.0, 4.6),
        ph_text(
            next(phid), "caption", ph_type="body", idx=11,
            x=0.5, y=5.25, w=10.2, h=1.75, sz=20, color=GRAPHITE, align="l",
        ),
        ph_text(
            next(phid), "timecode", ph_type="body", idx=12,
            x=10.9, y=6.75, w=1.9, h=0.45, sz=14, color=BURGUNDY, bold=True, align="r", anchor="ctr",
        ),
    ]

    # ingredients_list: heading, frame left, bulleted items right
    bodies["ingredients_list"] = [
        deco_rect(next(phid), 0, 0, SLIDE_W_IN, 0.14, BURGUNDY),
        ph_text(
            next(phid), "heading", ph_type="title", idx=None,
            x=0.5, y=0.35, w=9.0, h=0.95, sz=28, color=BURGUNDY, bold=True, align="l",
        ),
        ph_pic(next(phid), "frame", 10, 0.5, 1.5, 5.1, 5.5),
        ph_text(
            next(phid), "items", ph_type="body", idx=11,
            x=6.0, y=1.5, w=6.83, h=5.5, sz=16, color=GRAPHITE, align="l",
        ),
    ]

    # steps_grid: heading + 6 cell groups (number, timecode, frame, text)
    parts = [deco_rect(next(phid), 0, 0, SLIDE_W_IN, 0.14, BURGUNDY)]
    parts.append(
        ph_text(
            next(phid), "heading", ph_type="title", idx=None,
            x=0.5, y=0.3, w=9.5, h=0.9, sz=26, color=BURGUNDY, bold=True, align="l",
        )
    )
    for i, cell in enumerate(cell_specs(), 1):
        parts.append(
            ph_text(
                next(phid), f"cell{i}_number", ph_type="body", idx=10 + (i - 1) * 4,
                x=cell["number"][0], y=cell["number"][1], w=0.55, h=BADGE_H,
                sz=13, color=BURGUNDY, bold=True, align="l", anchor="ctr",
            )
        )
        parts.append(
            ph_text(
                next(phid), f"cell{i}_timecode", ph_type="body", idx=11 + (i - 1) * 4,
                x=cell["timecode"][0], y=cell["timecode"][1], w=1.3, h=BADGE_H,
                sz=11, color=GRAPHITE, align="r", anchor="ctr",
            )
        )
        parts.append(
            ph_pic(
                next(phid), f"cell{i}_frame", 12 + (i - 1) * 4,
                cell["frame"][0], cell["frame"][1], cell["frame"][2], cell["frame"][3],
            )
        )
        parts.append(
            ph_text(
                next(phid), f"cell{i}_text", ph_type="body", idx=13 + (i - 1) * 4,
                x=cell["text"][0], y=cell["text"][1], w=cell["text"][2], h=cell["text"][3],
                sz=13, color=GRAPHITE, align="l",
            )
        )
    bodies["steps_grid"] = parts

    # serving: heading, frame left, caption right, timecode bottom-right
    bodies["serving"] = [
        deco_rect(next(phid), 0, 0, SLIDE_W_IN, 0.14, BURGUNDY),
        ph_text(
            next(phid), "heading", ph_type="title", idx=None,
            x=0.5, y=0.35, w=9.0, h=0.95, sz=28, color=BURGUNDY, bold=True, align="l",
        ),
        ph_pic(next(phid), "frame", 10, 0.5, 1.5, 5.1, 5.5),
        ph_text(
            next(phid), "caption", ph_type="body", idx=11,
            x=6.0, y=1.5, w=6.83, h=4.6, sz=20, color=GRAPHITE, align="l",
        ),
        ph_text(
            next(phid), "timecode", ph_type="body", idx=12,
            x=11.18, y=6.5, w=1.65, h=0.5, sz=14, color=BURGUNDY, bold=True, align="r", anchor="ctr",
        ),
    ]
    return bodies


FOOTER_TEXT = "Razkadr"


def _strip_shapes(layout) -> None:
    """Remove every existing shape from a layout's spTree (the default
    template layouts are placeholder furniture we do not need)."""
    sp_tree = layout.shapes._spTree
    for tag in ("p:sp", "p:pic", "p:graphicFrame", "p:grpSp", "p:cxnSp"):
        for el in list(sp_tree.findall(qn(tag))):
            sp_tree.remove(el)


def _add_xml_shape(layout, xml: str) -> None:
    from pptx.oxml import parse_xml

    layout.shapes._spTree.append(parse_xml(xml))


def _brand_footer(layout) -> None:
    """Small static brand text bottom-left, baked into the layout."""
    _add_xml_shape(
        layout,
        static_text(
            999, "deco-footer",
            x=0.3, y=SLIDE_H_IN - 0.42, w=1.4, h=0.32,
            sz=9, color=CHAMPAGNE, bold=True,
        ),
    )


def build(out_path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Emu(_emu(SLIDE_W_IN))
    prs.slide_height = Emu(_emu(SLIDE_H_IN))
    master = prs.slide_masters[0]

    bases = list(master.slide_layouts)[:5]  # any five parts; contents replaced
    names = ["title", "hero_caption", "ingredients_list", "steps_grid", "serving"]
    bodies = layout_bodies()

    for base, name in zip(bases, names):
        _strip_shapes(base)
        base.name = name
        base.background.fill.solid()
        base.background.fill.fore_color.rgb = RGBColor.from_string(CREAM)
        for xml in bodies[name]:
            _add_xml_shape(base, xml)
        _brand_footer(base)

    # drop the six layouts we do not use so the template has exactly five
    keep_parts = {base.part for base in bases}
    id_lst = master.element.find(qn("p:sldLayoutIdLst"))
    for sld_layout_id in list(id_lst):
        r_id = sld_layout_id.get(qn("r:id"))
        if master.part.rels[r_id].target_part not in keep_parts:
            master.part.drop_rel(r_id)
            id_lst.remove(sld_layout_id)

    core = prs.core_properties
    core.title = "Razkadr default template"
    core.subject = "razkadr render reference template (naryad 022)"
    core.author = "Razkadr"
    core.comments = "Built by scripts/templates/build_default.py; do not edit by hand."
    core.category = "template"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "--out",
        default=str(Path(__file__).resolve().parents[2] / TEMPLATE_RELPATH),
        help=f"output path (default: repo {TEMPLATE_RELPATH})",
    )
    args = parser.parse_args()
    out = build(Path(args.out))
    print(f"template written: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
