"""Template check for user PPTX templates (naryad 023).

A user template follows the naming convention of docs/templates.md:
layouts are named ``razkadr:<макет>`` and slot placeholders ``slot:...``
(``slot:title``, ``slot:caption``, ``slot:image_1``, ``slot:timecode``,
...). The reference template is accepted too: its layouts carry the bare
macro names and its placeholders the internal slot names.

``check_template`` walks the five macro layouts and reports, per layout,
whether it exists and which required/optional slots are present or
missing. The template is renderable (exit 0) when every layout exists and
every required slot is in place; missing optional slots are notes, not
errors.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from razkadr.render.pptx import USER_LAYOUT_PREFIX

# (slot label, required, accepted placeholder names) per macro layout;
# label is the user-dialect name - it is what docs/templates.md teaches
SLOT_SPEC: dict[str, list[tuple[str, bool, tuple[str, ...]]]] = {
    "title": [
        ("slot:title", True, ("slot:title", "heading")),
        ("slot:subtitle", False, ("slot:subtitle", "subheading")),
        ("slot:image_1", True, ("slot:image_1", "frame")),
    ],
    "hero_caption": [
        ("slot:image_1", True, ("slot:image_1", "frame")),
        ("slot:caption", True, ("slot:caption", "caption")),
        ("slot:timecode", False, ("slot:timecode", "timecode")),
    ],
    "ingredients_list": [
        ("slot:title", True, ("slot:title", "heading")),
        ("slot:image_1", True, ("slot:image_1", "frame")),
        ("slot:list", True, ("slot:list", "items")),
    ],
    "steps_grid": [
        ("slot:title", True, ("slot:title", "heading")),
        # at least one complete cell group; groups 2..6 are optional
        ("slot:image_1", True, ("slot:image_1", "cell1_frame")),
        ("slot:text_1", True, ("slot:text_1", "cell1_text")),
        ("slot:number_1", False, ("slot:number_1", "cell1_number")),
        ("slot:timecode_1", False, ("slot:timecode_1", "cell1_timecode")),
        *[
            (f"slot:{kind}_{n}", False, (f"slot:{kind}_{n}", f"cell{n}_{k}"))
            for n in range(2, 7)
            for kind, k in (("image", "frame"), ("text", "text"), ("number", "number"), ("timecode", "timecode"))
        ],
    ],
    "serving": [
        ("slot:title", True, ("slot:title", "heading")),
        ("slot:image_1", True, ("slot:image_1", "frame")),
        ("slot:caption", True, ("slot:caption", "caption")),
        ("slot:timecode", False, ("slot:timecode", "timecode")),
    ],
}


class TemplateCheckError(Exception):
    """The template file cannot be inspected (not a PPTX)."""


def check_template(template: Path) -> tuple[bool, list[str]]:
    """Inspect the template; returns (renderable, report lines)."""
    try:
        prs = Presentation(str(template))
    except Exception as exc:  # zipfile.BadZipFile, ValueError, ...
        raise TemplateCheckError(f"not a readable PPTX: {exc}") from exc
    layouts = {layout.name: layout for layout in prs.slide_masters[0].slide_layouts}
    lines: list[str] = []
    ok = True
    for macro, slots in SLOT_SPEC.items():
        user_name = f"{USER_LAYOUT_PREFIX}{macro}"
        layout = layouts.get(user_name) or layouts.get(macro)
        if layout is None:
            ok = False
            lines.append(f"{macro}: layout missing (expected '{user_name}' or '{macro}')")
            continue
        found = {shape.name for shape in layout.shapes}
        dialect = "user" if layouts.get(user_name) else "reference"
        lines.append(f"{macro}: layout ok ({dialect} naming)")
        for label, required, accepted in slots:
            present = any(name in found for name in accepted)
            if present:
                lines.append(f"  {label}: ok")
            elif required:
                ok = False
                lines.append(f"  {label}: MISSING (required)")
            else:
                lines.append(f"  {label}: absent (optional)")
    return ok, lines
