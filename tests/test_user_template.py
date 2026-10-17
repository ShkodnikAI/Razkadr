"""User template tests (naryad 023).

Three generated templates drive the suite: the reference one (bare
layout names, internal slot names), a user-dialect one (``razkadr:*``
layouts, ``slot:*`` placeholders, renumbered ph idx, different geometry
and colors, 4 cell groups, no optional timecode slots) and a broken one
(missing the serving layout). template-check must report 0/0/1 and the
renderer must fill the user template respecting its own geometry.
"""

from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

import pytest
from pptx import Presentation

from razkadr.cli import EXIT_FILE_ERROR, EXIT_INVALID, EXIT_OK
from razkadr.cli import main as cli_main
from razkadr.render.pptx import RenderError, render_pptx
from razkadr.render.template_check import check_template

REPO = Path(__file__).parent.parent
EXAMPLES = REPO / "examples"
EMU_PER_IN = 914400


def _build_module():
    spec = importlib.util.spec_from_file_location(
        "build_default", REPO / "scripts/templates/build_default.py"
    )
    module = importlib.util.module_from_spec(spec)
    sys.modules["build_default"] = module
    spec.loader.exec_module(module)
    return module


def build_user_template(out: Path, *, with_serving: bool = True) -> Path:
    """A user-dialect template: razkadr:* layouts, slot:* placeholders,
    image left / text right (mirrored vs the reference), champagne
    headings, 4 cell groups, no timecode slots, ph idx renumbered 1..N."""
    bd = _build_module()
    from pptx import Presentation as P
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.util import Emu

    prs = P()
    prs.slide_width = Emu(bd._emu(13.333))
    prs.slide_height = Emu(bd._emu(7.5))
    master = prs.slide_masters[0]
    phid = iter(range(2, 300))

    ACCENT, TEXT = "B08830", "1F1F1F"  # champagne-ish accent, dark text

    bodies = {
        "razkadr:title": [
            bd.ph_text(next(phid), "slot:title", ph_type="title", idx=None, x=6.7, y=2.0, w=6.1, h=1.8, sz=36, color=ACCENT, bold=True),
            bd.ph_text(next(phid), "slot:subtitle", ph_type="body", idx=1, x=6.7, y=4.0, w=6.0, h=1.7, sz=18, color=TEXT),
            bd.ph_pic(next(phid), "slot:image_1", 2, 0.0, 0, 6.3, 7.5),  # left half
        ],
        "razkadr:hero_caption": [
            bd.ph_pic(next(phid), "slot:image_1", 1, 0.4, 0.4, 7.4, 6.7),  # tall left
            bd.ph_text(next(phid), "slot:caption", ph_type="body", idx=2, x=8.1, y=1.0, w=4.8, h=4.6, sz=18, color=TEXT),
            bd.deco_rect(next(phid), 0, 0, 13.333, 0.2, ACCENT),
        ],
        "razkadr:ingredients_list": [
            bd.ph_text(next(phid), "slot:title", ph_type="title", idx=None, x=0.5, y=0.3, w=9.0, h=0.9, sz=26, color=ACCENT, bold=True),
            bd.ph_pic(next(phid), "slot:image_1", 1, 7.6, 1.4, 5.2, 5.6),  # image right
            bd.ph_text(next(phid), "slot:list", ph_type="body", idx=2, x=0.5, y=1.4, w=6.7, h=5.6, sz=15, color=TEXT),
        ],
        "razkadr:steps_grid": [bd.ph_text(next(phid), "slot:title", ph_type="title", idx=None, x=0.5, y=0.3, w=9.5, h=0.9, sz=24, color=ACCENT, bold=True)],
        "razkadr:serving": [
            bd.ph_text(next(phid), "slot:title", ph_type="title", idx=None, x=0.5, y=0.3, w=9.0, h=0.9, sz=26, color=ACCENT, bold=True),
            bd.ph_pic(next(phid), "slot:image_1", 1, 0.5, 1.4, 6.0, 5.6),
            bd.ph_text(next(phid), "slot:caption", ph_type="body", idx=2, x=6.8, y=1.4, w=6.0, h=4.4, sz=18, color=TEXT),
        ],
    }
    # 4 complete cell groups (renumbered idx 1..), no timecode slots
    for n in range(1, 5):
        col, row = (n - 1) % 2, (n - 1) // 2
        cx, cy = 0.5 + col * 6.5, 1.4 + row * 3.0
        bodies["razkadr:steps_grid"] += [
            bd.ph_text(next(phid), f"slot:number_{n}", ph_type="body", idx=1 + (n - 1) * 3, x=cx, y=cy, w=0.5, h=0.4, sz=12, color=ACCENT, bold=True),
            bd.ph_pic(next(phid), f"slot:image_{n}", 2 + (n - 1) * 3, cx, cy + 0.45, 2.6, 2.3),
            bd.ph_text(next(phid), f"slot:text_{n}", ph_type="body", idx=3 + (n - 1) * 3, x=cx + 2.7, y=cy + 0.45, w=3.5, h=2.3, sz=12, color=TEXT),
        ]

    names = ["razkadr:title", "razkadr:hero_caption", "razkadr:ingredients_list", "razkadr:steps_grid"]
    if with_serving:
        names.append("razkadr:serving")
    bases = list(master.slide_layouts)[: len(names)]
    for base, name in zip(bases, names):
        sp_tree = base.shapes._spTree
        for tag in ("p:sp", "p:pic", "p:graphicFrame", "p:grpSp", "p:cxnSp"):
            for el in list(sp_tree.findall(qn(tag))):
                sp_tree.remove(el)
        base.name = name
        base.background.fill.solid()
        base.background.fill.fore_color.rgb = RGBColor.from_string("FBF7EF")
        for xml in bodies[name]:
            bd._add_xml_shape(base, xml)
    keep = {base.part for base in bases}
    id_lst = master.element.find(qn("p:sldLayoutIdLst"))
    for sld_layout_id in list(id_lst):
        r_id = sld_layout_id.get(qn("r:id"))
        if master.part.rels[r_id].target_part not in keep:
            master.part.drop_rel(r_id)
            id_lst.remove(sld_layout_id)
    prs.save(str(out))
    return out


def make_frames(slides_doc: dict, root: Path) -> None:
    from PIL import Image

    paths = []
    for slide in slides_doc["slides"]:
        slots = slide["slots"]
        if "frame" in slots:
            paths.append(slots["frame"]["path"])
        for cell in slots.get("cells", []):
            paths.append(cell["frame"]["path"])
    for i, rel in enumerate(sorted(set(paths))):
        p = root / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (640, 360), ((40 * i) % 256, 96, 160)).save(p, "JPEG", quality=85)


def example_slides(tmp_path: Path) -> tuple[dict, Path]:
    doc = json.loads((EXAMPLES / "artifacts/slides.json").read_text(encoding="utf-8"))
    (tmp_path / "slides.json").write_text(json.dumps(doc, ensure_ascii=False), encoding="utf-8")
    make_frames(doc, tmp_path)
    return doc, tmp_path


# --- template-check ---------------------------------------------------------------------


def test_check_reference_template_ok() -> None:
    ok, lines = check_template(REPO / "templates/razkadr-default.pptx")
    assert ok is True
    assert not any("MISSING" in line for line in lines)


def test_check_user_template_ok_with_optional_absent(tmp_path: Path) -> None:
    tpl = build_user_template(tmp_path / "user.pptx")
    ok, lines = check_template(tpl)
    assert ok is True, lines
    assert any("slot:timecode_1: absent (optional)" in line for line in lines)
    assert any("slot:image_5: absent (optional)" in line for line in lines)


def test_check_missing_layout_is_code_1(tmp_path: Path) -> None:
    tpl = build_user_template(tmp_path / "broken.pptx", with_serving=False)
    ok, lines = check_template(tpl)
    assert ok is False
    assert any("serving: layout missing" in line for line in lines)


def test_check_missing_required_slot_is_code_1(tmp_path: Path) -> None:
    """A user layout without the required image slot fails the check."""
    tpl = build_user_template(tmp_path / "incomplete.pptx")
    from pptx import Presentation as P

    prs = P(str(tpl))
    for lay in prs.slide_masters[0].slide_layouts:
        if lay.name == "razkadr:hero_caption":
            for shape in list(lay.shapes):
                if shape.name == "slot:image_1":
                    el = shape._element
                    el.getparent().remove(el)
    prs.save(str(tpl))
    ok, lines = check_template(tpl)
    assert ok is False
    assert any("slot:image_1: MISSING (required)" in line for line in lines)


# --- render on the user template ----------------------------------------------------------


def test_render_on_user_template_keeps_its_design(tmp_path: Path) -> None:
    tpl = build_user_template(tmp_path / "user.pptx")
    doc, root = example_slides(tmp_path)
    stats = render_pptx(doc, tmp_path / "deck.pptx", template=tpl, base_dir=root)
    assert stats["slides"] == 4
    assert stats["pictures"] == 4  # title + hero + ingredients + 1 cell
    assert stats["warnings"] == []
    prs = Presentation(str(tmp_path / "deck.pptx"))
    # title slide: image on the LEFT half (user design), heading text right
    title = prs.slides[0]
    pics = [s for s in title.shapes if s.shape_type == 13]
    assert len(pics) == 1 and pics[0].left < prs.slide_width / 2
    texts = " ".join(
        s.text_frame.text for s in title.shapes if s.has_text_frame
    )
    assert doc["slides"][0]["slots"]["heading"] in texts
    # user steps_grid: only the first cell group of the example (1 cell)
    steps = prs.slides[3]
    step_pics = [s for s in steps.shapes if s.shape_type == 13]
    assert len(step_pics) == 1
    names = [s.name for s in steps.shapes]
    for i in range(2, 5):  # unused cell groups dropped
        assert f"cell{i}_text" not in names


def test_render_user_steps_grid_more_cells_than_groups(tmp_path: Path) -> None:
    """6 cells in slides.json, only 4 complete groups in the template:
    cells 5-6 are dropped with a warning, nothing silent."""
    tpl = build_user_template(tmp_path / "user.pptx")
    doc, root = example_slides(tmp_path)
    slots = doc["slides"][3]["slots"]
    slots["columns"] = 3
    slots["cells"] = [
        {
            "frame": {"path": f"frames/f{i}.jpg", "t": 10.0 * i, "episode_id": "e3"},
            "number": i,
            "text": f"Шаг {i}",
            "timecode": f"{i}:00",
        }
        for i in range(1, 7)
    ]
    make_frames(doc, root)
    stats = render_pptx(doc, tmp_path / "deck.pptx", template=tpl, base_dir=root)
    assert stats["pictures"] == 7  # title + hero + ingredients + 4 placed cells
    assert any("cells_overflow" in w or "no complete" in w for w in stats["warnings"])
    prs = Presentation(str(tmp_path / "deck.pptx"))
    texts = " ".join(
        s.text_frame.text for s in prs.slides[3].shapes if s.has_text_frame
    )
    assert "Шаг 4" in texts
    assert "Шаг 5" not in texts


def test_render_missing_layout_is_clear_error(tmp_path: Path) -> None:
    tpl = build_user_template(tmp_path / "broken.pptx", with_serving=False)
    doc, root = example_slides(tmp_path)
    # the example has no serving slide - add one to hit the missing layout
    doc["slides"].append(
        {
            "layout": "serving",
            "source_refs": ["e3"],
            "warnings": [],
            "slots": {
                "heading": "Подача",
                "frame": {"path": "frames/serv.jpg", "t": 120.0, "episode_id": "e3"},
                "caption": "Готово",
                "timecode": "2:00",
            },
        }
    )
    make_frames(doc, root)
    with pytest.raises(RenderError, match="no layout for 'serving'"):
        render_pptx(doc, tmp_path / "deck.pptx", template=tpl, base_dir=root)


# --- CLI -----------------------------------------------------------------------------------


def test_cli_template_check_reference_exit_0(tmp_path: Path, capsys) -> None:
    code = cli_main(["template-check", str(REPO / "templates/razkadr-default.pptx")])
    assert code == EXIT_OK
    assert "ready for render" in capsys.readouterr().out


def test_cli_template_check_broken_exit_1(tmp_path: Path, capsys) -> None:
    tpl = build_user_template(tmp_path / "broken.pptx", with_serving=False)
    code = cli_main(["template-check", str(tpl)])
    assert code == EXIT_INVALID
    err = capsys.readouterr().err
    assert "not ready" in err and "docs/templates.md" in err


def test_cli_template_check_missing_file_exit_2(tmp_path: Path, capsys) -> None:
    code = cli_main(["template-check", str(tmp_path / "nope.pptx")])
    assert code == EXIT_FILE_ERROR
    assert "cannot read" in capsys.readouterr().err


def test_cli_render_with_user_template(tmp_path: Path) -> None:
    tpl = build_user_template(tmp_path / "user.pptx")
    example_slides(tmp_path)
    code = cli_main(["render", str(tmp_path / "slides.json"), "--template", str(tpl), "--out", str(tmp_path / "deck.pptx")])
    assert code == EXIT_OK
    assert (tmp_path / "deck.pptx").is_file()
